import pytest
import os
import tempfile
from pathlib import Path
from docx import Document
from unittest.mock import MagicMock
from app.services.resume_service import ResumeService, ResumeExtractionError
from app.models.schemas import CandidateProfile, Project

@pytest.fixture
def resume_service(mocker):
    # Mock the LLM to avoid real API calls
    mocker.patch('app.services.resume_service.ChatGroq')
    service = ResumeService()
    # Mock the structured output chain invoke method
    service.structured_llm = MagicMock()
    return service

def test_extract_invalid_extension(resume_service):
    with pytest.raises(ResumeExtractionError, match="Unsupported file format"):
        # We also mock os.path.exists so it doesn't fail on "File not found"
        with pytest.MonkeyPatch.context() as m:
            m.setattr(os.path, 'exists', lambda x: True)
            resume_service.extract_text_from_file("invalid_resume.txt")

def test_legacy_doc_is_not_advertised_as_supported(resume_service):
    with pytest.raises(ResumeExtractionError, match="Unsupported file format"):
        with pytest.MonkeyPatch.context() as m:
            m.setattr(os.path, 'exists', lambda x: True)
            resume_service.extract_text_from_file("legacy_resume.doc")

def test_docx_extraction_includes_table_cells_in_document_order(resume_service):
    with tempfile.TemporaryDirectory(dir=Path.cwd()) as folder:
        path = Path(folder) / "resume.docx"
        document = Document()
        document.add_paragraph("Experience")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Engineer"
        table.cell(0, 1).text = "Built services"
        document.add_paragraph("Education")
        document.save(path)

        text = resume_service.extract_text_from_file(str(path))
        assert text.index("Experience") < text.index("Engineer") < text.index("Education")
        assert "Built services" in text

def test_available_sample_resume_preserves_its_role_and_experience_text(resume_service):
    sample_path = Path(__file__).resolve().parents[1] / "mock_resume.docx"
    text = resume_service.extract_text_from_file(str(sample_path))

    assert "Senior Backend Engineer at TechCorp" in text
    assert "5 years of experience" in text
    assert "Built scalable microservices" in text
    assert "Projects" not in text


def test_profile_generation_receives_the_extracted_resume_text(resume_service, mocker):
    sample_path = Path(__file__).resolve().parents[1] / "mock_resume.docx"
    text = resume_service.extract_text_from_file(str(sample_path))
    chain = MagicMock()
    chain.invoke.return_value = CandidateProfile(
        experience=["Senior Backend Engineer at TechCorp"], projects=[],
    )
    mocker.patch(
        "app.services.resume_service.ChatPromptTemplate.from_messages",
        return_value=MagicMock(__or__=lambda self, other: chain),
    )

    profile = resume_service.analyze_resume(text)

    assert profile.experience == ["Senior Backend Engineer at TechCorp"]
    assert profile.projects == []
    assert chain.invoke.call_args.args[0]["resume_text"] == text

def test_extract_missing_file(resume_service):
    with pytest.raises(FileNotFoundError, match="Resume file not found"):
        resume_service.extract_text_from_file("does_not_exist.pdf")

def test_analyze_empty_resume(resume_service):
    with pytest.raises(ResumeExtractionError, match="Extracted text is too short"):
        resume_service.analyze_resume("Too short")

def test_analyze_resume_success(resume_service, mocker):
    mock_profile = CandidateProfile(
        name="John Doe",
        education=["BSc Computer Science"],
        skills=["Python", "SQL"],
        projects=[Project(name="TestProj", description="Test", technologies=["Python"])],
        resume_summary="Good candidate",
        inferred_strengths=["Strong in backend"],
        potential_weaknesses=["Lacks frontend experience"],
        interviewable_technologies=["Python"]
    )
    
    # Create a mock chain
    mock_chain = MagicMock()
    mock_chain.invoke.return_value = mock_profile
    
    # We patch the | operator by mocking the prompt | llm result
    # In ResumeService, chain is created inside analyze_resume, so we patch the structured_llm invoke instead if possible.
    # Wait, chain.invoke is called.
    # Let's mock the ChatPromptTemplate.from_messages
    mocker.patch('app.services.resume_service.ChatPromptTemplate.from_messages', return_value=MagicMock(__or__=lambda self, other: mock_chain))
    
    result = resume_service.analyze_resume("This is a valid resume text with enough length to bypass the short check. Python SQL BSc Computer Science.")
    
    assert result.name == "John Doe"
    assert "Python" in result.skills
    assert result.resume_summary == "Good candidate"

def test_analyze_resume_llm_failure(resume_service, mocker):
    mock_chain = MagicMock()
    mock_chain.invoke.side_effect = Exception("LLM Error")
    mocker.patch('app.services.resume_service.ChatPromptTemplate.from_messages', return_value=MagicMock(__or__=lambda self, other: mock_chain))
    
    with pytest.raises(ResumeExtractionError, match="Failed to analyze resume with LLM"):
        resume_service.analyze_resume("This is a valid resume text with enough length to bypass the short check. Python SQL BSc Computer Science.")

